from io import BytesIO

import docx
import pymupdf
import pytest
from openpyxl import Workbook

from app.documents.extract import ExtractionError, detect_kind, extract_pages


def test_detect_kind() -> None:
    assert detect_kind("rapport.PDF", None) == "pdf"
    assert detect_kind("x", "application/pdf") == "pdf"
    assert detect_kind("budget.xlsx", "application/octet-stream") == "xlsx"
    assert detect_kind("photo.jpg", "image/jpeg") is None


def test_pdf_pages() -> None:
    pdf = pymupdf.open()
    for text in ("Objectif général du projet", "Résultat 1.1 : AVEC fonctionnelles"):
        pdf.new_page().insert_text((72, 72), text)
    pages = extract_pages(pdf.tobytes(), "pdf")
    assert [p.number for p in pages] == [1, 2]
    assert "AVEC fonctionnelles" in pages[1].text


def test_docx_paragraphs_and_tables() -> None:
    document = docx.Document()
    document.add_paragraph("Activité A1.1.1 : Former les membres des AVEC")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Kits"
    table.rows[0].cells[1].text = "3 600 USD"
    buffer = BytesIO()
    document.save(buffer)
    text = extract_pages(buffer.getvalue(), "docx")[0].text
    assert "Former les membres" in text
    assert "Kits | 3 600 USD" in text


def test_xlsx_one_page_per_sheet() -> None:
    wb = Workbook()
    wb.active.title = "Budget"  # type: ignore[union-attr]
    wb.active.append(["Ligne", "Montant"])  # type: ignore[union-attr]
    wb.create_sheet("Indicateurs").append(["I1", 600])
    buffer = BytesIO()
    wb.save(buffer)
    pages = extract_pages(buffer.getvalue(), "xlsx")
    assert pages[0].text.startswith("[Feuille : Budget]")
    assert "I1 | 600" in pages[1].text


def test_long_text_is_paginated() -> None:
    text = "\n".join(f"Paragraphe {i} " + "x" * 200 for i in range(60))
    pages = extract_pages(text.encode(), "txt")
    assert len(pages) > 3
    assert "Paragraphe 0" in pages[0].text


def test_empty_document_is_rejected() -> None:
    with pytest.raises(ExtractionError):
        extract_pages(b"   \n  ", "txt")

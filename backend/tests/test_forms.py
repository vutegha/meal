from io import BytesIO
from typing import Any
from uuid import uuid4

from httpx import AsyncClient
from openpyxl import load_workbook

from tests.conftest import register
from tests.test_organizations import add_member, login
from tests.test_projects import create_project

PDM: dict[str, Any] = {
    "title": "Suivi post-distribution des kits",
    "fields": [
        {"key": "menage", "label": "Taille du ménage", "type": "integer", "required": True},
        {"key": "recu", "label": "Kit reçu complet ?", "type": "yesno", "required": True},
        {
            "key": "usage",
            "label": "Usage principal",
            "type": "select",
            "options": ["Consommation", "Vente", "Autre"],
        },
        {
            "key": "problemes",
            "label": "Problèmes rencontrés",
            "type": "multiselect",
            "options": ["Distance", "Attente", "Qualité"],
        },
        {"key": "date", "label": "Date de distribution", "type": "date"},
        {"key": "remarque", "label": "Remarques", "type": "text"},
    ],
}


async def setup(client: AsyncClient) -> tuple[dict[str, Any], str]:
    ctx = await register(client, "pm@forms.org", "ONG Formulaires")
    project = await create_project(client, ctx)
    return ctx, f"/api/v1/orgs/{ctx['org_id']}/projects/{project}/forms"


async def test_form_lifecycle_and_summary(client: AsyncClient) -> None:
    ctx, base = await setup(client)
    h = ctx["headers"]
    bad = {**PDM, "fields": [{"key": "x", "label": "Choix", "type": "select", "options": ["A"]}]}
    assert (await client.post(base, headers=h, json=bad)).status_code == 422
    form = (await client.post(base, headers=h, json=PDM)).json()
    assert form["status"] == "draft"
    url = f"{base}/{form['id']}"

    answer = {"menage": "6", "recu": True, "usage": "Vente", "problemes": ["Attente", "Distance"]}
    r = await client.post(f"{url}/submissions", headers=h, json={"answers": answer})
    assert r.status_code == 409  # brouillon : pas encore de réponses

    await client.patch(url, headers=h, json={"status": "published"})
    await add_member(client, ctx, "agent@forms.org", "field_agent")
    agent = await login(client, "agent@forms.org")
    client_uuid = str(uuid4())
    body = {"answers": answer, "location": "Kiwanja", "client_uuid": client_uuid}
    r = await client.post(f"{url}/submissions", headers=agent, json=body)
    assert r.status_code == 201, r.text
    saved = r.json()
    # Valeurs normalisées : entier, choix dans l'ordre du formulaire.
    assert saved["answers"] == {
        "menage": 6,
        "recu": True,
        "usage": "Vente",
        "problemes": ["Distance", "Attente"],
    }
    again = await client.post(f"{url}/submissions", headers=agent, json=body)
    assert again.status_code == 200 and again.json()["id"] == saved["id"]

    for answers, error in (
        ({"recu": False}, "« Taille du ménage » : réponse obligatoire"),
        ({"menage": 2.5, "recu": True}, "« Taille du ménage » : un nombre entier est attendu"),
        ({"menage": 2, "recu": True, "usage": "Don"}, "« Usage principal » : choix inconnu"),
        ({"menage": 2, "recu": True, "inconnue": 1}, "Questions inconnues : inconnue"),
        ({"menage": 2, "recu": True, "date": "31/03/2026"}, "« Date de distribution »"),
    ):
        r = await client.post(f"{url}/submissions", headers=agent, json={"answers": answers})
        assert r.status_code == 422 and error in r.json()["detail"], r.text
    await client.post(
        f"{url}/submissions",
        headers=agent,
        json={
            "answers": {
                "menage": 4,
                "recu": False,
                "usage": "Consommation",
                "remarque": "Riz abîmé",
            }
        },
    )

    summary = (await client.get(f"{url}/summary", headers=h)).json()
    assert summary["submissions"] == 2
    by_key = {f["key"]: f for f in summary["fields"]}
    assert by_key["menage"]["mean"] == 5 and by_key["menage"]["total"] == 10
    assert by_key["recu"]["counts"] == {"yes": 1, "no": 1}
    assert by_key["usage"]["counts"] == {"Consommation": 1, "Vente": 1, "Autre": 0}
    assert by_key["problemes"]["counts"]["Attente"] == 1
    assert by_key["remarque"]["samples"] == ["Riz abîmé"]

    # Les questions sont figées dès la première réponse.
    r = await client.patch(url, headers=h, json={"fields": PDM["fields"][:2]})
    assert r.status_code == 409
    listed = (await client.get(base, headers=h)).json()
    assert listed[0]["submissions"] == 2

    rows = (await client.get(f"{url}/submissions", headers=h)).json()
    assert rows[0]["submitter_name"] == "Agent"
    assert (await client.get(f"{url}/submissions", headers=agent)).status_code == 403

    export = await client.get(f"{url}/export.xlsx", headers=h)
    sheet = load_workbook(BytesIO(export.content)).active
    assert sheet is not None
    assert [c.value for c in sheet[1]][5:7] == ["Taille du ménage", "Kit reçu complet ?"]
    assert [c.value for c in sheet[2]][5:9] == [6, "Oui", "Vente", "Distance, Attente"]
    audit = (await client.get(f"/api/v1/orgs/{ctx['org_id']}/audit", headers=h)).json()
    assert [e["data"]["rows"] for e in audit if e["action"] == "form.exported"] == [len(rows)]

    await client.patch(url, headers=h, json={"status": "closed"})
    r = await client.post(f"{url}/submissions", headers=agent, json={"answers": answer})
    assert r.status_code == 409


async def test_forms_are_isolated(client: AsyncClient) -> None:
    ctx, base = await setup(client)
    form = (await client.post(base, headers=ctx["headers"], json=PDM)).json()
    other = await register(client, "x@other.org", "Autre ONG")
    r = await client.get(f"{base}/{form['id']}", headers=other["headers"])
    assert r.status_code == 404


async def test_personal_questions_stay_out_of_summary_and_export(client: AsyncClient) -> None:
    ctx, base = await setup(client)
    h = ctx["headers"]
    fields = [
        {"key": "nom", "label": "Nom du chef de ménage", "type": "text", "personal": True},
        {"key": "age", "label": "Âge", "type": "integer", "personal": True},
        {"key": "handicap", "label": "Handicap", "type": "yesno", "personal": True},
        {"key": "remarque", "label": "Remarques", "type": "text"},
    ]
    form = (await client.post(base, headers=h, json={"title": "Enquête", "fields": fields})).json()
    url = f"{base}/{form['id']}"
    await client.patch(url, headers=h, json={"status": "published"})
    for nom, age in (("Amani Furaha", 71), ("Neema Kahindo", 34)):
        answers = {"nom": nom, "age": age, "handicap": age > 60, "remarque": "RAS"}
        r = await client.post(f"{url}/submissions", headers=h, json={"answers": answers})
        assert r.status_code == 201, r.text

    by_key = {f["key"]: f for f in (await client.get(f"{url}/summary", headers=h)).json()["fields"]}
    assert by_key["nom"]["personal"] and by_key["nom"]["samples"] == []
    assert by_key["nom"]["answered"] == 2
    assert by_key["age"]["max"] is None and by_key["age"]["mean"] is None
    # Un effectif reste utile pour ventiler les résultats.
    assert by_key["handicap"]["counts"] == {"yes": 1, "no": 1}
    assert by_key["remarque"]["samples"] == ["RAS", "RAS"]

    sheet = load_workbook(BytesIO((await client.get(f"{url}/export.xlsx", headers=h)).content))
    assert sheet.active is not None
    assert [c.value for c in sheet.active[1]][5:] == ["Remarques"]
    full = await client.get(f"{url}/export.xlsx", headers=h, params={"personal": "true"})
    sheet = load_workbook(BytesIO(full.content))
    assert sheet.active is not None
    assert "Amani Furaha" in [c.value for c in sheet.active[2]]
    audit = (await client.get(f"/api/v1/orgs/{ctx['org_id']}/audit", headers=h)).json()
    exported = [e["data"]["personal"] for e in audit if e["action"] == "form.exported"]
    assert sorted(exported) == [False, True]

    # Après les premières réponses, on peut encore marquer une question comme personnelle.
    fields[3]["personal"] = True
    assert (await client.patch(url, headers=h, json={"fields": fields})).status_code == 200
    fields[3]["label"] = "Autres remarques"
    assert (await client.patch(url, headers=h, json={"fields": fields})).status_code == 409

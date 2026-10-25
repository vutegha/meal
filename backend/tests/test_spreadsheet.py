from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.documents.spreadsheet import save_as_text


def test_text_starting_with_equals_is_not_a_formula() -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(['=HYPERLINK("https://exemple.test","Cliquez")', "+243 99", 12])
    sheet = load_workbook(BytesIO(save_as_text(wb))).active
    assert sheet is not None
    cell = sheet["A1"]
    assert cell.data_type == "s" and str(cell.value).startswith("=HYPERLINK")
    assert sheet["B1"].value == "+243 99" and sheet["C1"].value == 12

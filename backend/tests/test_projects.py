from io import BytesIO
from typing import Any

from httpx import AsyncClient
from openpyxl import load_workbook

from tests.conftest import register
from tests.test_organizations import add_member, login


async def create_project(client: AsyncClient, ctx: dict[str, Any], code: str = "P1") -> str:
    r = await client.post(
        f"/api/v1/orgs/{ctx['org_id']}/projects",
        headers=ctx["headers"],
        json={
            "code": code,
            "title": "Résilience communautaire au Nord-Kivu",
            "donor": "UE",
            "currency": "EUR",
            "start_date": "2026-01-01",
            "end_date": "2027-12-31",
            "zones": ["Goma", "Rutshuru"],
        },
    )
    assert r.status_code == 201, r.text
    return str(r.json()["id"])


async def add_node(
    client: AsyncClient,
    ctx: dict[str, Any],
    base: str,
    level: str,
    title: str,
    parent_id: str | None = None,
    code: str = "",
) -> Any:
    r = await client.post(
        f"{base}/logframe/nodes",
        headers=ctx["headers"],
        json={"level": level, "title": title, "parent_id": parent_id, "code": code},
    )
    assert r.status_code == 201, r.text
    return r.json()


async def build_logframe(client: AsyncClient, ctx: dict[str, Any], base: str) -> dict[str, Any]:
    goal = await add_node(client, ctx, base, "goal", "Réduire la vulnérabilité", code="OG")
    outcome = await add_node(
        client, ctx, base, "outcome", "Les ménages diversifient leurs revenus", goal["id"], "OS1"
    )
    output = await add_node(
        client, ctx, base, "output", "Des AVEC sont fonctionnelles", outcome["id"], "R1.1"
    )
    activity = await add_node(
        client, ctx, base, "activity", "Former 30 AVEC", output["id"], "A1.1.1"
    )
    return {"goal": goal, "outcome": outcome, "output": output, "activity": activity}


async def test_project_crud_and_unique_code(client: AsyncClient) -> None:
    ctx = await register(client, "pm@example.org")
    pid = await create_project(client, ctx)
    base = f"/api/v1/orgs/{ctx['org_id']}/projects"

    projects = (await client.get(base, headers=ctx["headers"])).json()
    assert [p["code"] for p in projects] == ["P1"]
    assert projects[0]["zones"] == ["Goma", "Rutshuru"]

    dup = await client.post(base, headers=ctx["headers"], json={"code": "P1", "title": "Autre"})
    assert dup.status_code == 409

    other = await create_project(client, ctx, code="P2")
    r = await client.patch(f"{base}/{other}", headers=ctx["headers"], json={"code": "P1"})
    assert r.status_code == 409

    r = await client.patch(f"{base}/{pid}", headers=ctx["headers"], json={"status": "active"})
    assert r.json()["status"] == "active"
    bad = await client.patch(
        f"{base}/{pid}", headers=ctx["headers"], json={"end_date": "2025-01-01"}
    )
    assert bad.status_code == 422

    assert (await client.delete(f"{base}/{pid}", headers=ctx["headers"])).status_code == 204
    assert (await client.get(f"{base}/{pid}", headers=ctx["headers"])).status_code == 404


async def test_logframe_hierarchy_and_tree(client: AsyncClient) -> None:
    ctx = await register(client, "pm@example.org")
    base = f"/api/v1/orgs/{ctx['org_id']}/projects/{await create_project(client, ctx)}"
    nodes = await build_logframe(client, ctx, base)

    # Une activité ne peut pas être rattachée directement à un objectif général.
    r = await client.post(
        f"{base}/logframe/nodes",
        headers=ctx["headers"],
        json={"level": "activity", "title": "X", "parent_id": nodes["goal"]["id"]},
    )
    assert r.status_code == 422
    # Un effet doit avoir un parent.
    r = await client.post(
        f"{base}/logframe/nodes", headers=ctx["headers"], json={"level": "outcome", "title": "X"}
    )
    assert r.status_code == 422

    second = await add_node(
        client, ctx, base, "activity", "Suivre les AVEC", nodes["output"]["id"], "A1.1.2"
    )
    assert second["position"] == 1

    tree = (await client.get(f"{base}/logframe", headers=ctx["headers"])).json()
    output = tree[0]["children"][0]["children"][0]
    assert [a["code"] for a in output["children"]] == ["A1.1.1", "A1.1.2"]

    r = await client.patch(
        f"{base}/logframe/nodes/{second['id']}",
        headers=ctx["headers"],
        json={"title": "Accompagner les AVEC", "assumptions": "Accès sécurisé"},
    )
    assert r.json()["title"] == "Accompagner les AVEC"

    # Supprimer le résultat supprime ses activités.
    r = await client.delete(
        f"{base}/logframe/nodes/{nodes['output']['id']}", headers=ctx["headers"]
    )
    assert r.status_code == 204
    tree = (await client.get(f"{base}/logframe", headers=ctx["headers"])).json()
    assert tree[0]["children"][0]["children"] == []


async def test_budget_lines_expenses_and_summary(client: AsyncClient) -> None:
    ctx = await register(client, "pm@example.org")
    base = f"/api/v1/orgs/{ctx['org_id']}/projects/{await create_project(client, ctx)}"
    nodes = await build_logframe(client, ctx, base)

    r = await client.post(
        f"{base}/budget/lines",
        headers=ctx["headers"],
        json={"label": "X", "quantity": 1, "unit_cost": 1, "activity_id": nodes["output"]["id"]},
    )
    assert r.status_code == 422  # rattachement à un résultat interdit

    training = (
        await client.post(
            f"{base}/budget/lines",
            headers=ctx["headers"],
            json={
                "activity_id": nodes["activity"]["id"],
                "label": "Kits de formation",
                "donor_line_code": "3.1",
                "quantity": 30,
                "unit": "kit",
                "unit_cost": "45.50",
                "frequency": 2,
            },
        )
    ).json()
    assert training["planned"] == "2730.00"
    support = (
        await client.post(
            f"{base}/budget/lines",
            headers=ctx["headers"],
            json={"label": "Loyer bureau", "quantity": 12, "unit_cost": 100},
        )
    ).json()

    for amount in ("1000", "500.25"):
        r = await client.post(
            f"{base}/budget/lines/{training['id']}/expenses",
            headers=ctx["headers"],
            json={"amount": amount, "spent_on": "2026-03-01", "reference": "FAC-1"},
        )
        assert r.status_code == 201
    r = await client.post(
        f"{base}/budget/lines/{support['id']}/expenses",
        headers=ctx["headers"],
        json={"amount": 1300, "spent_on": "2026-03-01"},
    )

    lines = (await client.get(f"{base}/budget/lines", headers=ctx["headers"])).json()
    assert {line["label"]: line["spent"] for line in lines} == {
        "Kits de formation": "1500.25",
        "Loyer bureau": "1300.00",
    }

    summary = (await client.get(f"{base}/budget/summary", headers=ctx["headers"])).json()
    assert summary["currency"] == "EUR"
    assert summary["planned"] == "3930.00"
    assert summary["spent"] == "2800.25"
    assert summary["over_budget_lines"] == 1
    by_title = {a["title"]: a for a in summary["by_activity"]}
    assert by_title["Former 30 AVEC"]["execution_rate"] == round(1500.25 / 2730, 4)
    assert by_title["Coûts de support"]["planned"] == "1200.00"

    r = await client.patch(
        f"{base}/budget/lines/{training['id']}", headers=ctx["headers"], json={"quantity": 40}
    )
    assert r.json()["planned"] == "3640.00"
    assert r.json()["spent"] == "1500.25"


async def test_indicators_achievement_and_checks(client: AsyncClient) -> None:
    ctx = await register(client, "pm@example.org")
    base = f"/api/v1/orgs/{ctx['org_id']}/projects/{await create_project(client, ctx)}"
    nodes = await build_logframe(client, ctx, base)

    trained = (
        await client.post(
            f"{base}/indicators",
            headers=ctx["headers"],
            json={
                "node_id": nodes["output"]["id"],
                "code": "I1",
                "name": "Nombre de membres d'AVEC formés",
                "unit": "personnes",
                "target": 600,
                "disaggregations": ["sexe"],
                "source_of_verification": "Listes de présence",
            },
        )
    ).json()
    income = (
        await client.post(
            f"{base}/indicators",
            headers=ctx["headers"],
            json={
                "node_id": nodes["outcome"]["id"],
                "code": "I2",
                "name": "Revenu mensuel moyen des ménages",
                "unit": "USD",
                "baseline": 40,
                "target": 60,
                "aggregation": "latest",
            },
        )
    ).json()
    assert trained["achieved"] is None

    async def record(indicator: dict[str, Any], value: float, end: str, **extra: Any) -> Any:
        return await client.post(
            f"{base}/indicators/{indicator['id']}/values",
            headers=ctx["headers"],
            json={"period_start": "2026-01-01", "period_end": end, "value": value, **extra},
        )

    assert (await record(trained, 200, "2026-03-31")).status_code == 201
    assert (await record(trained, 100, "2026-06-30")).status_code == 201
    # Une ligne désagrégée n'entre pas dans le total.
    r = await record(trained, 120, "2026-06-30", disaggregation={"sexe": "femmes"})
    assert r.status_code == 201
    r = await record(trained, 1, "2026-06-30", disaggregation={"age": "18-25"})
    assert r.status_code == 422
    assert (await record(income, 45, "2026-03-31")).status_code == 201
    assert (await record(income, 50, "2026-06-30")).status_code == 201

    indicators = {
        i["code"]: i
        for i in (await client.get(f"{base}/indicators", headers=ctx["headers"])).json()
    }
    assert indicators["I1"]["achieved"] == "300.0000"
    assert indicators["I1"]["achievement_rate"] == 0.5
    assert indicators["I2"]["achieved"] == "50.0000"
    assert indicators["I2"]["achievement_rate"] == 0.5

    check = (await client.get(f"{base}/logframe/check", headers=ctx["headers"])).json()
    assert [n["code"] for n in check["nodes_without_indicator"]] == ["OG"]
    assert [i["code"] for i in check["indicators_without_source"]] == ["I2"]
    assert [n["code"] for n in check["activities_without_budget"]] == ["A1.1.1"]

    r = await client.patch(
        f"{base}/indicators/{income['id']}",
        headers=ctx["headers"],
        json={"source_of_verification": "Enquête ménages", "target": None},
    )
    assert r.json()["source_of_verification"] == "Enquête ménages"
    assert r.json()["achievement_rate"] is None


async def test_export_xlsx(client: AsyncClient) -> None:
    ctx = await register(client, "pm@example.org")
    base = f"/api/v1/orgs/{ctx['org_id']}/projects/{await create_project(client, ctx)}"
    nodes = await build_logframe(client, ctx, base)
    await client.post(
        f"{base}/indicators",
        headers=ctx["headers"],
        json={"node_id": nodes["output"]["id"], "code": "I1", "name": "AVEC fonctionnelles"},
    )
    await client.post(
        f"{base}/budget/lines",
        headers=ctx["headers"],
        json={
            "activity_id": nodes["activity"]["id"],
            "label": "Kits",
            "quantity": 2,
            "unit_cost": 10,
        },
    )
    r = await client.get(f"{base}/export.xlsx", headers=ctx["headers"])
    assert r.status_code == 200
    assert "cadre-logique-P1.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Cadre logique", "Budget", "Indicateurs"]
    logframe = [row for row in wb["Cadre logique"].iter_rows(min_row=2, values_only=True)]
    assert [row[1] for row in logframe] == ["OG", "OS1", "R1.1", "A1.1.1"]
    assert logframe[2][3] == "I1 AVEC fonctionnelles"
    assert list(wb["Budget"].iter_rows(min_row=2, values_only=True))[0][8] == 20


async def test_project_roles(client: AsyncClient) -> None:
    admin = await register(client, "admin@example.org")
    base = f"/api/v1/orgs/{admin['org_id']}/projects/{await create_project(client, admin)}"
    nodes = await build_logframe(client, admin, base)
    indicator = (
        await client.post(
            f"{base}/indicators",
            headers=admin["headers"],
            json={"node_id": nodes["output"]["id"], "name": "AVEC fonctionnelles"},
        )
    ).json()
    await add_member(client, admin, "terrain@example.org", "field_agent")
    await add_member(client, admin, "finance@example.org", "finance")
    field = await login(client, "terrain@example.org")
    finance = await login(client, "finance@example.org")

    # L'agent terrain lit le projet et saisit des valeurs, mais ne planifie pas.
    assert (await client.get(f"{base}/logframe", headers=field)).status_code == 200
    r = await client.post(
        f"{base}/indicators/{indicator['id']}/values",
        headers=field,
        json={"period_start": "2026-01-01", "period_end": "2026-01-31", "value": 3},
    )
    assert r.status_code == 201
    r = await client.post(
        f"{base}/logframe/nodes", headers=field, json={"level": "goal", "title": "X"}
    )
    assert r.status_code == 403
    r = await client.post(
        f"{base}/budget/lines", headers=field, json={"label": "X", "quantity": 1, "unit_cost": 1}
    )
    assert r.status_code == 403

    # La finance gère le budget mais pas le cadre logique.
    r = await client.post(
        f"{base}/budget/lines", headers=finance, json={"label": "X", "quantity": 1, "unit_cost": 1}
    )
    assert r.status_code == 201
    r = await client.post(
        f"{base}/logframe/nodes", headers=finance, json={"level": "goal", "title": "X"}
    )
    assert r.status_code == 403
    r = await client.post(
        f"/api/v1/orgs/{admin['org_id']}/projects",
        headers=finance,
        json={"code": "P2", "title": "Autre projet"},
    )
    assert r.status_code == 403


async def test_projects_are_isolated(client: AsyncClient) -> None:
    a = await register(client, "a@example.org", org="Org A")
    b = await register(client, "b@example.org", org="Org B")
    base_a = f"/api/v1/orgs/{a['org_id']}/projects/{await create_project(client, a)}"
    nodes_a = await build_logframe(client, a, base_a)
    project_b = await create_project(client, b)
    base_b = f"/api/v1/orgs/{b['org_id']}/projects/{project_b}"

    # B ne voit pas le projet de A, ni via l'organisation A ni en substituant l'identifiant.
    assert (await client.get(base_a, headers=b["headers"])).status_code == 404
    project_a = base_a.rsplit("/", 1)[1]
    r = await client.get(
        f"/api/v1/orgs/{b['org_id']}/projects/{project_a}/logframe", headers=b["headers"]
    )
    assert r.status_code == 404

    # B ne peut pas rattacher ses éléments à un nœud du projet de A.
    r = await client.post(
        f"{base_b}/logframe/nodes",
        headers=b["headers"],
        json={"level": "outcome", "title": "X", "parent_id": nodes_a["goal"]["id"]},
    )
    assert r.status_code == 404
    r = await client.post(
        f"{base_b}/indicators",
        headers=b["headers"],
        json={"node_id": nodes_a["output"]["id"], "name": "Intrus"},
    )
    assert r.status_code == 404
    r = await client.post(
        f"{base_b}/budget/lines",
        headers=b["headers"],
        json={
            "activity_id": nodes_a["activity"]["id"],
            "label": "X",
            "quantity": 1,
            "unit_cost": 1,
        },
    )
    assert r.status_code == 404
    # Ni désigner comme responsable quelqu'un d'une autre organisation.
    goal_b = await add_node(client, b, base_b, "goal", "Objectif B")
    me_a = (await client.get("/api/v1/auth/me", headers=a["headers"])).json()
    r = await client.post(
        f"{base_b}/indicators",
        headers=b["headers"],
        json={"node_id": goal_b["id"], "name": "Indicateur B", "owner_id": me_a["id"]},
    )
    assert r.status_code == 422

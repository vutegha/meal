from io import BytesIO
from typing import Any

import docx
import pymupdf
from httpx import AsyncClient

from tests.conftest import register
from tests.fake_llm import FakeLLM
from tests.test_organizations import add_member, login
from tests.test_periodic import MARCH
from tests.test_tor import fake_llm, setup  # noqa: F401

__all__ = ["fake_llm"]

TOR_TEMPLATE: dict[str, Any] = {
    "kind": "tor",
    "name": "TdR format UE",
    "donor": "ue",
    "sections": [
        {"key": "contexte", "title": "Contexte", "guidance": "Deux paragraphes au plus."},
        {"key": "budget", "title": "Budget détaillé"},
        {"key": "visibilite", "title": "Visibilité du bailleur", "guidance": "Logo UE."},
    ],
    "layout": {"header": "ONG Kivu · Confidentiel", "footer": "Siège : Goma", "color": "#1d4ed8"},
}


async def test_builtin_sections(client: AsyncClient) -> None:
    ctx = await register(client, "m@example.org")
    r = await client.get(f"/api/v1/orgs/{ctx['org_id']}/templates/builtin", headers=ctx["headers"])
    builtin = r.json()
    assert [s["key"] for s in builtin["tor"] if s["computed"]] == ["budget"]
    assert {s["key"] for s in builtin["periodic"] if s["computed"]} == {
        "activites",
        "indicateurs",
        "budget",
    }


async def test_donor_template_drives_tor(client: AsyncClient, fake_llm: FakeLLM) -> None:
    ctx, base, nodes = await setup(client)
    h, templates = ctx["headers"], f"/api/v1/orgs/{ctx['org_id']}/templates"
    bad = {**TOR_TEMPLATE, "sections": [TOR_TEMPLATE["sections"][0]] * 2}
    assert (await client.post(templates, headers=h, json=bad)).status_code == 422
    r = await client.post(templates, headers=h, json=TOR_TEMPLATE)
    assert r.status_code == 201, r.text
    template = r.json()

    # Le projet du bailleur « UE » prend le modèle UE (casse ignorée).
    activity = nodes["activity"]["id"]
    tor = (await client.post(f"{base}/activities/{activity}/tor", headers=h)).json()
    assert tor["template_id"] == template["id"]
    assert [s["title"] for s in tor["sections"]] == [
        "Contexte",
        "Budget détaillé",
        "Visibilité du bailleur",
    ]
    assert "Formateurs" in tor["sections"][1]["content"]

    r = await client.post(f"{base}/activities/{activity}/tor/generate", headers=h, json={})
    job = (
        await client.get(f"/api/v1/orgs/{ctx['org_id']}/jobs/{r.json()['id']}", headers=h)
    ).json()
    assert job["status"] == "succeeded", job["error"]
    ask = fake_llm.calls[-1]["content"][1]["text"]
    assert "- contexte : Contexte (consigne de l'organisation : Deux paragraphes au plus.)" in ask
    assert "visibilite" in ask and "budget" not in ask
    tor = (await client.get(f"{base}/tors/{tor['id']}", headers=h)).json()
    contents = {s["key"]: s["content"] for s in tor["sections"]}
    assert contents["contexte"].startswith("Les ménages ciblés")
    assert "Formateurs" in contents["budget"]
    assert contents["visibilite"] == "[À compléter]"

    # La mise en page du modèle s'applique aux exports.
    pdf = await client.get(f"{base}/tors/{tor['id']}/export.pdf", headers=h)
    text = pymupdf.open(stream=pdf.content, filetype="pdf")[0].get_text()
    assert "ONG Kivu · Confidentiel" in text and "Page 1 / " in text
    word = docx.Document(
        BytesIO((await client.get(f"{base}/tors/{tor['id']}/export.docx", headers=h)).content)
    )
    assert word.sections[0].header.paragraphs[0].text == "ONG Kivu · Confidentiel"
    assert word.sections[0].footer.paragraphs[0].text.startswith("Siège : Goma")


async def test_org_default_and_single_default(client: AsyncClient) -> None:
    ctx, base, _ = await setup(client)
    h, templates = ctx["headers"], f"/api/v1/orgs/{ctx['org_id']}/templates"
    periodic = {
        "kind": "periodic",
        "name": "Rapport trimestriel interne",
        "is_default": True,
        "sections": [
            {"key": "resume", "title": "Faits marquants"},
            {"key": "budget", "title": "Budget"},
        ],
    }
    first = (await client.post(templates, headers=h, json=periodic)).json()
    second = (
        await client.post(templates, headers=h, json={**periodic, "name": "Autre modèle"})
    ).json()
    listed = {t["id"]: t["is_default"] for t in (await client.get(templates, headers=h)).json()}
    assert listed == {first["id"]: False, second["id"]: True}

    report = (await client.post(f"{base}/periodic-reports", headers=h, json=MARCH)).json()
    assert report["template_id"] == second["id"]
    assert [s["key"] for s in report["sections"]] == ["resume", "budget"]
    assert report["sections"][1]["content"].startswith("|")  # tableau calculé

    # Choix explicite d'un modèle pour un document.
    chosen = (
        await client.post(
            f"{base}/periodic-reports", headers=h, json={**MARCH, "template_id": first["id"]}
        )
    ).json()
    assert chosen["template_id"] == first["id"]

    # Modèle supprimé : les documents gardent leurs sections.
    assert (await client.delete(f"{templates}/{second['id']}", headers=h)).status_code == 204
    kept = (await client.get(f"{base}/periodic-reports/{report['id']}", headers=h)).json()
    assert kept["template_id"] is None and len(kept["sections"]) == 2


async def test_template_roles_and_isolation(client: AsyncClient) -> None:
    ctx = await register(client, "admin@example.org")
    templates = f"/api/v1/orgs/{ctx['org_id']}/templates"
    created = (await client.post(templates, headers=ctx["headers"], json=TOR_TEMPLATE)).json()
    await add_member(client, ctx, "meal@example.org", "meal_officer")
    meal = await login(client, "meal@example.org")
    assert len((await client.get(templates, headers=meal)).json()) == 1
    assert (await client.post(templates, headers=meal, json=TOR_TEMPLATE)).status_code == 403

    other = await register(client, "other@example.org", "Autre ONG")
    r = await client.patch(
        f"/api/v1/orgs/{other['org_id']}/templates/{created['id']}",
        headers=other["headers"],
        json={"name": "Piraté"},
    )
    assert r.status_code == 404
    r = await client.patch(
        f"{templates}/{created['id']}", headers=ctx["headers"], json={"name": "TdR UE 2026"}
    )
    assert r.json()["name"] == "TdR UE 2026"

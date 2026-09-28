# PyMuPDF n'annote pas les types de son API Story/DocumentWriter.
# mypy: disable-error-code="no-untyped-call"
"""Rendu des documents rédigés (TdR, rapports) en Word et en PDF.

Le contenu est saisi en Markdown simple : titres (#), paragraphes, listes (- ou 1.), tableaux
(| a | b |) et gras (**texte**). Un seul analyseur alimente les deux rendus, pour que le Word
et le PDF disent la même chose.
"""

import html
import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import Literal

BlockKind = Literal["heading", "paragraph", "bullets", "numbered", "table"]


@dataclass
class Block:
    kind: BlockKind
    text: str = ""
    level: int = 0
    items: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


@dataclass
class Section:
    title: str
    content: str


@dataclass
class Figure:
    """Image JPEG ou PNG avec sa légende (les deux formats lus par Word et par le PDF)."""

    data: bytes
    caption: str


@dataclass
class Layout:
    """Mise en page d'un modèle de l'organisation : en-tête, pied de page et couleur des titres."""

    header: str = ""
    footer: str = ""
    color: str = "#0f5b52"


@dataclass
class RenderedDocument:
    title: str
    subtitle: str
    meta: list[tuple[str, str]]
    sections: list[Section]
    figures: list[Figure] = field(default_factory=list)
    figures_title: str = "Photos"
    layout: Layout = field(default_factory=Layout)


_BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_HEADING = re.compile(r"^(#{1,3})\s+(.*)$")
_SEPARATOR = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$")
_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def parse(markdown: str) -> list[Block]:
    blocks: list[Block] = []
    paragraph: list[str] = []

    def close_paragraph() -> None:
        if paragraph:
            blocks.append(Block("paragraph", text=" ".join(paragraph)))
            paragraph.clear()

    for raw in markdown.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            close_paragraph()
            continue
        if heading := _HEADING.match(line):
            close_paragraph()
            blocks.append(Block("heading", text=heading[2].strip(), level=len(heading[1])))
        elif line.lstrip().startswith("|"):
            close_paragraph()
            if _SEPARATOR.match(line.strip()):
                continue
            if blocks and blocks[-1].kind == "table":
                blocks[-1].rows.append(_cells(line))
            else:
                blocks.append(Block("table", rows=[_cells(line)]))
        elif (item := _BULLET.match(line)) or (item := _NUMBERED.match(line)):
            close_paragraph()
            kind: BlockKind = "bullets" if _BULLET.match(line) else "numbered"
            if blocks and blocks[-1].kind == kind:
                blocks[-1].items.append(item[1].strip())
            else:
                blocks.append(Block(kind, items=[item[1].strip()]))
        else:
            paragraph.append(line.strip())
    close_paragraph()
    return blocks


# --- Word ------------------------------------------------------------------------


def _runs(paragraph: object, text: str) -> None:
    """Ajoute le texte au paragraphe python-docx en respectant le **gras**."""
    for index, part in enumerate(_BOLD.split(text)):
        if part:
            paragraph.add_run(part).bold = index % 2 == 1  # type: ignore[attr-defined]


def _apply_layout(word: object, layout: Layout) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    color = RGBColor.from_string(layout.color.lstrip("#").upper())
    for name in ("Title", "Heading 1", "Heading 2"):
        word.styles[name].font.color.rgb = color  # type: ignore[attr-defined]
    section = word.sections[0]  # type: ignore[attr-defined]
    if layout.header:
        header = section.header.paragraphs[0]
        header.text = layout.header
        header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.text = f"{layout.footer}    " if layout.footer else ""
    footer.add_run("Page ")
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    for run in footer.runs:
        run.font.size = Pt(8)


def to_docx(document: RenderedDocument) -> bytes:
    import docx
    from docx.shared import Pt

    word = docx.Document()
    style = word.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    _apply_layout(word, document.layout)
    word.add_heading(document.title, level=0)
    if document.subtitle:
        word.add_paragraph(document.subtitle).italic = True  # type: ignore[attr-defined]
    if document.meta:
        table = word.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        for label, value in document.meta:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = label, value
            cells[0].paragraphs[0].runs[0].bold = True

    for number, section in enumerate(document.sections, start=1):
        word.add_heading(f"{number}. {section.title}", level=1)
        for block in parse(section.content):
            if block.kind == "heading":
                word.add_heading(block.text, level=min(block.level + 1, 4))
            elif block.kind == "paragraph":
                _runs(word.add_paragraph(), block.text)
            elif block.kind in ("bullets", "numbered"):
                list_style = "List Bullet" if block.kind == "bullets" else "List Number"
                for item in block.items:
                    _runs(word.add_paragraph(style=list_style), item)
            else:
                width = max(len(row) for row in block.rows)
                table = word.add_table(rows=0, cols=width)
                table.style = "Table Grid"
                for row_index, row in enumerate(block.rows):
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row, strict=False):
                        _runs(cell.paragraphs[0], value)
                        if row_index == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
    if document.figures:
        from docx.shared import Cm

        word.add_heading(f"{len(document.sections) + 1}. {document.figures_title}", level=1)
        for figure in document.figures:
            word.add_picture(BytesIO(figure.data), width=Cm(12))
            if figure.caption:
                word.add_paragraph(figure.caption).runs[0].italic = True
    buffer = BytesIO()
    word.save(buffer)
    return buffer.getvalue()


# --- PDF -------------------------------------------------------------------------


def _inline(text: str) -> str:
    return _BOLD.sub(r"<b>\1</b>", html.escape(text))


def to_html(document: RenderedDocument) -> str:
    parts = [f"<h1>{html.escape(document.title)}</h1>"]
    if document.subtitle:
        parts.append(f"<p class='subtitle'>{html.escape(document.subtitle)}</p>")
    if document.meta:
        rows = "".join(
            f"<tr><th>{html.escape(label)}</th><td>{html.escape(value)}</td></tr>"
            for label, value in document.meta
        )
        parts.append(f"<table class='meta'>{rows}</table>")
    for number, section in enumerate(document.sections, start=1):
        parts.append(f"<h2>{number}. {html.escape(section.title)}</h2>")
        for block in parse(section.content):
            if block.kind == "heading":
                parts.append(f"<h3>{_inline(block.text)}</h3>")
            elif block.kind == "paragraph":
                parts.append(f"<p>{_inline(block.text)}</p>")
            elif block.kind in ("bullets", "numbered"):
                tag = "ul" if block.kind == "bullets" else "ol"
                items = "".join(f"<li>{_inline(item)}</li>" for item in block.items)
                parts.append(f"<{tag}>{items}</{tag}>")
            else:
                head, *body = block.rows
                header = "".join(f"<th>{_inline(cell)}</th>" for cell in head)
                rows = "".join(
                    "<tr>" + "".join(f"<td>{_inline(cell)}</td>" for cell in row) + "</tr>"
                    for row in body
                )
                parts.append(f"<table><tr>{header}</tr>{rows}</table>")
    if document.figures:
        parts.append(
            f"<h2>{len(document.sections) + 1}. {html.escape(document.figures_title)}</h2>"
        )
        for index, figure in enumerate(document.figures):
            parts.append(f"<p><img src='figure-{index}' width='380'/></p>")
            if figure.caption:
                parts.append(f"<p class='caption'>{html.escape(figure.caption)}</p>")
    return "\n".join(parts)


_CSS = """
body { font-family: sans-serif; font-size: 10.5pt; line-height: 1.35; }
h1 { font-size: 18pt; color: #0f5b52; margin-bottom: 4pt; }
h2 { font-size: 13pt; color: #0f5b52; margin-top: 14pt; }
h3 { font-size: 11pt; }
.subtitle { font-style: italic; color: #555; }
.caption { font-style: italic; color: #555; font-size: 9pt; }
table { border-collapse: collapse; width: 100%; margin: 6pt 0; }
th, td { border: 0.5pt solid #999; padding: 3pt 5pt; text-align: left; vertical-align: top; }
th { background-color: #eef4f3; }
"""


def to_pdf(document: RenderedDocument) -> bytes:
    import pymupdf

    archive = pymupdf.Archive()
    for index, figure in enumerate(document.figures):
        archive.add(figure.data, f"figure-{index}")
    css = _CSS.replace("#0f5b52", document.layout.color)
    story = pymupdf.Story(html=to_html(document), user_css=css, archive=archive)
    buffer = BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    page = pymupdf.paper_rect("a4")
    area = page + (56, 56, -56, -56)
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(area)
        story.draw(device)
        writer.end_page()
    writer.close()
    return _stamp(buffer.getvalue(), document.layout)


def _stamp(data: bytes, layout: Layout) -> bytes:
    """En-tête, pied de page et numéro de page sur chaque page."""
    import pymupdf

    pdf = pymupdf.open(stream=data, filetype="pdf")
    grey = (0.4, 0.4, 0.4)
    for number in range(1, pdf.page_count + 1):
        page = pdf[number - 1]
        width, height = page.rect.width, page.rect.height
        if layout.header:
            page.insert_textbox(
                pymupdf.Rect(56, 22, width - 56, 48), layout.header, fontsize=8, color=grey
            )
        footer = f"{layout.footer}    " if layout.footer else ""
        page.insert_textbox(
            pymupdf.Rect(56, height - 40, width - 56, height - 16),
            f"{footer}Page {number} / {pdf.page_count}",
            fontsize=8,
            color=grey,
            align=pymupdf.TEXT_ALIGN_RIGHT,
        )
    return bytes(pdf.tobytes(garbage=3, deflate=True))

"""Extraction du texte des documents, page par page (ou feuille par feuille)."""

from dataclasses import dataclass
from io import BytesIO

SUPPORTED_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "text/plain": "txt",
    "text/markdown": "txt",
}

EXTENSIONS = {".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx", ".txt": "txt", ".md": "txt"}

# Nombre approximatif de caractères par « page » pour les formats sans pagination.
CHARS_PER_PAGE = 3000


class ExtractionError(Exception):
    pass


@dataclass
class Page:
    number: int
    text: str


def detect_kind(filename: str, content_type: str | None) -> str | None:
    if content_type in SUPPORTED_TYPES:
        return SUPPORTED_TYPES[content_type]
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return EXTENSIONS.get(suffix)


def _paginate(text: str) -> list[Page]:
    """Découpe un texte continu en pages virtuelles, aux sauts de paragraphe."""
    pages: list[Page] = []
    current: list[str] = []
    size = 0
    for paragraph in text.split("\n"):
        if size + len(paragraph) > CHARS_PER_PAGE and current:
            pages.append(Page(len(pages) + 1, "\n".join(current)))
            current, size = [], 0
        current.append(paragraph)
        size += len(paragraph) + 1
    if any(line.strip() for line in current):
        pages.append(Page(len(pages) + 1, "\n".join(current)))
    return pages


def _pdf(data: bytes) -> list[Page]:
    import pymupdf

    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:  # type: ignore[no-untyped-call]
            return [Page(i + 1, page.get_text()) for i, page in enumerate(doc)]
    except Exception as exc:  # pymupdf lève des exceptions génériques
        raise ExtractionError("PDF illisible") from exc


def _docx(data: bytes) -> list[Page]:
    import docx

    try:
        document = docx.Document(BytesIO(data))
    except Exception as exc:
        raise ExtractionError("Document Word illisible") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return _paginate("\n".join(lines))


def _xlsx(data: bytes) -> list[Page]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError("Classeur Excel illisible") from exc
    pages = []
    for index, sheet in enumerate(workbook.worksheets):
        rows = [
            " | ".join("" if v is None else str(v) for v in row)
            for row in sheet.iter_rows(values_only=True)
            if any(v is not None for v in row)
        ]
        pages.append(Page(index + 1, f"[Feuille : {sheet.title}]\n" + "\n".join(rows)))
    return pages


def _txt(data: bytes) -> list[Page]:
    for encoding in ("utf-8", "cp1252"):
        try:
            return _paginate(data.decode(encoding))
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Encodage du texte non reconnu")


def extract_pages(data: bytes, kind: str) -> list[Page]:
    pages = {"pdf": _pdf, "docx": _docx, "xlsx": _xlsx, "txt": _txt}[kind](data)
    if not any(p.text.strip() for p in pages):
        raise ExtractionError(
            "Aucun texte trouvé. Le document est peut-être scanné (la reconnaissance de "
            "caractères n'est pas encore disponible)."
        )
    return pages

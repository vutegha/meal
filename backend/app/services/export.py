"""Export Excel, PDF et Word du cadre logique, du budget et des indicateurs."""

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.documents.render import RenderedDocument, Section
from app.models import BudgetLine, Project
from app.schemas.project import NodeTree
from app.services import projects as svc
from app.services.periodic import day
from app.services.tor import fmt

LEVEL_LABELS = {
    "goal": "Objectif général",
    "outcome": "Objectif spécifique / effet",
    "output": "Résultat / extrant",
    "activity": "Activité",
    "sub_activity": "Sous-activité",
}

HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(ws: Worksheet, headers: list[str], widths: list[int]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, width in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + index)].width = width
    ws.freeze_panes = "A2"


def _wrap(ws: Worksheet) -> None:
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


async def logframe_workbook(session: AsyncSession, project: Project) -> bytes:
    nodes = await svc.list_nodes(session, project.id)
    indicators = await svc.list_indicators(session, project.id)
    by_node: dict[object, list[str]] = {}
    sources: dict[object, list[str]] = {}
    for indicator in indicators:
        by_node.setdefault(indicator.node_id, []).append(
            f"{indicator.code} {indicator.name}".strip()
        )
        if indicator.source_of_verification:
            sources.setdefault(indicator.node_id, []).append(indicator.source_of_verification)

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Cadre logique"
    _sheet(
        ws,
        ["Niveau", "Code", "Intitulé", "Indicateurs", "Sources de vérification", "Hypothèses"],
        [26, 10, 60, 50, 40, 40],
    )

    def walk(tree: list[NodeTree], depth: int) -> None:
        for node in tree:
            ws.append(
                [
                    LEVEL_LABELS[node.level.value],
                    node.code,
                    "    " * depth + node.title,
                    "\n".join(by_node.get(node.id, [])),
                    "\n".join(sources.get(node.id, [])),
                    node.assumptions,
                ]
            )
            if depth == 0:
                ws.cell(ws.max_row, 3).font = Font(bold=True)
            walk(node.children, depth + 1)

    walk(svc.build_tree(nodes), 0)
    _wrap(ws)

    codes = {n.id: f"{n.code} {n.title}".strip() for n in nodes}
    budget = wb.create_sheet("Budget")
    _sheet(
        budget,
        [
            "Code bailleur",
            "Activité",
            "Libellé",
            "Catégorie",
            "Quantité",
            "Unité",
            "Coût unitaire",
            "Fréquence",
            f"Prévu ({project.currency})",
            f"Dépensé ({project.currency})",
            "Estimation",
        ],
        [14, 40, 40, 18, 10, 10, 14, 10, 16, 16, 11],
    )
    spent = await svc.spent_by_line(session, project.id)
    lines = await session.scalars(
        select(BudgetLine)
        .where(BudgetLine.project_id == project.id)
        .order_by(BudgetLine.donor_line_code, BudgetLine.created_at)
    )
    for line in lines:
        budget.append(
            [
                line.donor_line_code,
                codes[line.activity_id] if line.activity_id else "Coûts de support",
                line.label,
                line.category,
                float(line.quantity),
                line.unit,
                float(line.unit_cost),
                float(line.frequency),
                float(line.planned),
                float(spent.get(line.id, 0)),
                "oui" if line.is_estimate else "",
            ]
        )
    _wrap(budget)

    sheet = wb.create_sheet("Indicateurs")
    _sheet(
        sheet,
        [
            "Code",
            "Indicateur",
            "Rattaché à",
            "Unité",
            "Référence",
            "Cible",
            "Atteint",
            "Taux",
            "Désagrégations",
            "Source",
            "Méthode",
            "Fréquence",
        ],
        [10, 50, 40, 10, 11, 11, 11, 9, 24, 30, 30, 14],
    )
    for indicator in indicators:
        out = svc.indicator_out(indicator)
        sheet.append(
            [
                out.code,
                out.name,
                codes.get(out.node_id, ""),
                out.unit,
                float(out.baseline) if out.baseline is not None else None,
                float(out.target) if out.target is not None else None,
                float(out.achieved) if out.achieved is not None else None,
                out.achievement_rate,
                ", ".join(out.disaggregations),
                out.source_of_verification,
                out.collection_method,
                out.frequency,
            ]
        )
        sheet.cell(sheet.max_row, 8).number_format = "0%"
    _wrap(sheet)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- PDF (et Word) ------------------------------------------------------------------------


def _cell(text: str) -> str:
    """Texte sûr pour une cellule de tableau Markdown."""
    return " ".join(text.replace("|", "/").split()) or "–"


def _table(headers: list[str], rows: list[list[str]]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    lines += ["| " + " | ".join(_cell(c) for c in row) + " |" for row in rows]
    return "\n".join(lines)


def _percent(value: float | None) -> str:
    return f"{round(value * 100)} %" if value is not None else "–"


def _number(value: Decimal | None) -> str:
    return fmt(value) if value is not None else "–"


async def logframe_document(session: AsyncSession, project: Project) -> RenderedDocument:
    """Cadre logique, indicateurs (avec cibles par période) et budget, en un document."""
    nodes = await svc.list_nodes(session, project.id)
    indicators = await svc.list_indicators(session, project.id)
    by_node: dict[object, list[str]] = {}
    for indicator in indicators:
        by_node.setdefault(indicator.node_id, []).append(
            f"{indicator.code} {indicator.name}".strip()
        )

    rows: list[list[str]] = []

    def walk(tree: list[NodeTree], depth: int) -> None:
        for node in tree:
            rows.append(
                [
                    LEVEL_LABELS[node.level.value],
                    node.code,
                    "– " * depth + node.title,
                    " ; ".join(by_node.get(node.id, [])),
                    node.assumptions,
                ]
            )
            walk(node.children, depth + 1)

    walk(svc.build_tree(nodes), 0)
    logframe = _table(["Niveau", "Code", "Intitulé", "Indicateurs", "Hypothèses"], rows)

    outs = [svc.indicator_out(i) for i in indicators]
    parts = [
        _table(
            ["Code", "Indicateur", "Référence", "Cible", "Atteint", "Taux"],
            [
                [
                    o.code,
                    f"{o.name} ({o.unit})" if o.unit else o.name,
                    _number(o.baseline),
                    _number(o.target),
                    _number(o.achieved),
                    _percent(o.achievement_rate),
                ]
                for o in outs
            ],
        )
    ]
    for o in outs:
        if o.period_targets:
            parts.append(f"### Cibles par période : {o.code} {o.name}".rstrip())
            parts.append(
                _table(
                    ["Période", "Cible", "Atteint", "Taux"],
                    [
                        [
                            f"{day(t.period_start)} – {day(t.period_end)}",
                            _number(t.target),
                            _number(t.achieved),
                            _percent(t.achievement_rate),
                        ]
                        for t in o.period_targets
                    ],
                )
            )

    summary = await svc.budget_summary(session, project)
    c = project.currency
    budget = [
        _table(
            ["Activité", f"Prévu ({c})", f"Dépensé ({c})", "Exécution"],
            [
                [a.title if not a.code else f"{a.code} {a.title}", fmt(a.planned), fmt(a.spent)]
                + [_percent(a.execution_rate)]
                for a in summary.by_activity
            ]
            + [
                ["**Total**", fmt(summary.planned), fmt(summary.spent)]
                + [_percent(summary.execution_rate)]
            ],
        )
    ]
    if project.exchange_rates:
        budget.append("### Taux de change")
        budget.append(
            _table(
                ["Devise", f"Valeur en {c}", "À partir du"],
                [
                    [
                        r["currency"],
                        fmt(Decimal(str(r["rate"])), 6).rstrip("0").rstrip(","),
                        day(date.fromisoformat(str(r["valid_from"]))),
                    ]
                    for r in project.exchange_rates
                ],
            )
        )

    dates = (
        f"{day(project.start_date)} – {day(project.end_date)}"
        if project.start_date and project.end_date
        else ""
    )
    meta = [
        (label, value)
        for label, value in (
            ("Bailleur", project.donor),
            ("Période", dates),
            ("Zones", ", ".join(project.zones)),
            ("Devise", c),
            ("Édité le", day(date.today())),
        )
        if value
    ]
    return RenderedDocument(
        title=f"{project.code} · {project.title}",
        subtitle="Cadre logique, indicateurs et budget",
        meta=meta,
        sections=[
            Section("Cadre logique", logframe),
            Section("Indicateurs", "\n\n".join(parts)),
            Section("Budget", "\n\n".join(budget)),
        ],
    )
